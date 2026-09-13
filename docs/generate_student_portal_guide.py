from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Preformatted,
    PageBreak,
)


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
DOCX_PATH = OUT_DIR / "Student_Portal_Microservices_Docker_Guide.docx"
PDF_PATH = OUT_DIR / "Student_Portal_Microservices_Docker_Guide.pdf"


TITLE = "Student Portal: Microservices, Docker, Database, and Frontend Guide"
SUBTITLE = "Beginner-friendly short guide for running and understanding the project"


SECTIONS = [
    (
        "1. What This Project Contains",
        [
            "Frontend: HTML, CSS, and JavaScript pages inside frontend/.",
            "Backend: multiple Spring Boot REST microservices inside backend/Admin/.",
            "Database: MySQL databases used by each service. Hibernate creates/updates tables from entities.",
            "Docker: docker-compose.yml runs MySQL, all backend services, and an Nginx frontend server.",
            "Portal_Service: new service for student attendance summaries and internal in-app mail.",
        ],
    ),
    (
        "2. Simple Architecture Flow",
        [
            "Admin creates master data: students, faculty, courses, timetable, fee details.",
            "Faculty logs in and sees courses assigned by admin.",
            "Faculty marks attendance by section and subject.",
            "Portal_Service stores attendance records in MySQL.",
            "Student opens Attendance and sees only their own subject-wise percentage.",
            "Mail is internal application mail only: sender chooses recipient role, name, and ID. Only that recipient can see the message inside the application.",
        ],
    ),
    (
        "3. Microservices Used",
        [
            "Admin_Register - admin login/register service on port 9093.",
            "Admin_Students - stores students created by admin on port 9098.",
            "Admin_Faculty - stores faculty created by admin on port 9102.",
            "Admin_Student_Course - stores student course data on port 9094.",
            "Admin_Faculty_Course - stores faculty course mapping on port 9095.",
            "Admin_Students_Timetable - student timetable service on port 9100.",
            "Admin_Faculty_Timetable - faculty timetable service on port 9099.",
            "AdminDashboard - aggregates counts from students/faculty services on port 9083.",
            "FacultyRegister - faculty login service on port 9070.",
            "FacultyHome - faculty course dashboard service on port 9071.",
            "MyFee - fee service on port 9096.",
            "Portal_Service - attendance and internal mail service on port 9110.",
        ],
    ),
    (
        "4. Important New APIs",
        [
            "POST /portal/attendance/mark - faculty submits attendance for a subject/date/section.",
            "GET /portal/attendance/student/{registerNumber}/summary - student sees subject-wise attendance.",
            "POST /portal/mail/send - sends mail inside the application only.",
            "GET /portal/mail/inbox?role=STUDENT&id=REGNO - loads messages for one exact recipient.",
        ],
    ),
    (
        "5. How Docker Combines Everything",
        [
            "docker-compose.yml starts one MySQL container first.",
            "The MySQL init script creates all required databases automatically.",
            "Each Spring Boot service receives its database URL using SPRING_DATASOURCE_URL.",
            "Backend services communicate by Docker service name, for example http://admin-students:9098.",
            "The frontend runs through Nginx on port 8080 and the browser calls localhost backend ports.",
        ],
    ),
    (
        "6. Run With Docker",
        [
            "Install Docker Desktop and keep it running.",
            "Open PowerShell in C:\\Htmlprojects\\Student_Portal.",
            "Run: docker compose up --build",
            "Open frontend: http://localhost:8080/FirstPage.html",
            "Stop all containers: press Ctrl+C, then run docker compose down.",
        ],
    ),
    (
        "7. Run Without Docker",
        [
            "Start MySQL locally.",
            "Create databases from docker/mysql/init/01-create-databases.sql or database/*.sql.",
            "Run start-all.bat to start all Spring Boot services in separate windows.",
            "Open frontend/FirstPage.html directly in the browser, or serve frontend/ with a small static server.",
        ],
    ),
    (
        "8. Database Notes",
        [
            "For Docker, no manual database creation is needed. The init SQL creates databases.",
            "For local MySQL, create the same databases before starting services.",
            "Hibernate ddl-auto=update lets Spring Boot create/update tables from entity classes.",
            "Database passwords are not hard-coded now. They come from environment variables.",
        ],
    ),
    (
        "9. Frontend Responsiveness Work",
        [
            "The existing design was kept: same colors, sidebar, buttons, cards, and page structure.",
            "Fixed fixed-margin sidebar label alignment by allowing text and icons to flex naturally.",
            "Added mobile/tablet breakpoints for forms, tables, mail, attendance grid, and cards.",
            "Tables scroll horizontally on small screens instead of breaking layout.",
            "Added smooth hover/focus transitions without changing the visual theme.",
        ],
    ),
]


CODE_BLOCKS = [
    (
        "Docker command",
        "cd C:\\Htmlprojects\\Student_Portal\n"
        "docker compose up --build\n"
        "open http://localhost:8080/FirstPage.html",
    ),
    (
        "Attendance request example",
        "POST http://localhost:9110/portal/attendance/mark\n"
        "{\n"
        '  "courseName": "DevOps & Microservices",\n'
        '  "section": "MCA",\n'
        '  "facultyName": "Faculty Name",\n'
        '  "facultyRegisterNumber": "FAC001",\n'
        '  "attendanceDate": "2026-07-17",\n'
        '  "students": [\n'
        '    {"studentName": "Student One", "studentRegisterNumber": "MCA001", "status": "PRESENT"}\n'
        "  ]\n"
        "}",
    ),
    (
        "Internal mail request example",
        "POST http://localhost:9110/portal/mail/send\n"
        "{\n"
        '  "senderRole": "FACULTY",\n'
        '  "senderName": "Faculty Name",\n'
        '  "senderId": "FAC001",\n'
        '  "recipientRole": "STUDENT",\n'
        '  "recipientName": "Student One",\n'
        '  "recipientId": "MCA001",\n'
        '  "subject": "Class update",\n'
        '  "body": "Tomorrow class starts at 9 AM."\n'
        "}",
    ),
]


PORT_TABLE = [
    ["Service", "Port", "Purpose"],
    ["frontend", "8080", "Nginx static frontend"],
    ["mysql", "3306", "MySQL database"],
    ["Portal_Service", "9110", "Attendance and in-app mail"],
    ["FacultyRegister", "9070", "Faculty login"],
    ["FacultyHome", "9071", "Faculty dashboard/courses"],
    ["Admin_Students", "9098", "Admin student data"],
    ["Admin_Faculty", "9102", "Admin faculty data"],
    ["AdminDashboard", "9083", "Dashboard counts"],
]


def add_docx_heading(document, text, level):
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Arial"
    return heading


def build_docx():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    title = document.add_heading(TITLE, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(SUBTITLE)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_docx_heading(document, "Fast Mental Model", 1)
    document.add_paragraph(
        "Think of the project as one frontend talking to many small Spring Boot services. "
        "Each service owns a focused job and stores data in MySQL. Docker starts all of them together."
    )

    add_docx_heading(document, "Core Ports", 1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for index, cell in enumerate(table.rows[0].cells):
        cell.text = PORT_TABLE[0][index]
    for row in PORT_TABLE[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value

    for heading, bullets in SECTIONS:
        add_docx_heading(document, heading, 1)
        for bullet in bullets:
            document.add_paragraph(bullet, style="List Bullet")

    add_docx_heading(document, "Simple Code Examples", 1)
    for heading, code in CODE_BLOCKS:
        add_docx_heading(document, heading, 2)
        paragraph = document.add_paragraph()
        run = paragraph.add_run(code)
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    add_docx_heading(document, "Learning Path", 1)
    for item in [
        "First understand one flow: faculty marks attendance, student sees summary.",
        "Then learn how REST endpoints move JSON between frontend and backend.",
        "Then learn Hibernate: entity class becomes database table.",
        "Finally learn Docker: compose file starts MySQL, services, and frontend together.",
    ]:
        document.add_paragraph(item, style="List Number")

    document.save(DOCX_PATH)


def build_pdf():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SmallCode", fontName="Courier", fontSize=8, leading=10))
    styles["Title"].alignment = 1
    styles["Heading1"].spaceBefore = 12
    styles["Heading1"].spaceAfter = 6

    story = [
        Paragraph(TITLE, styles["Title"]),
        Paragraph(SUBTITLE, styles["Normal"]),
        Spacer(1, 0.2 * inch),
        Paragraph("Fast Mental Model", styles["Heading1"]),
        Paragraph(
            "One frontend talks to many Spring Boot REST microservices. Each service owns a focused job and stores data in MySQL. Docker starts everything together.",
            styles["Normal"],
        ),
        Spacer(1, 0.15 * inch),
        Paragraph("Core Ports", styles["Heading1"]),
    ]

    table = Table(PORT_TABLE, hAlign="LEFT", colWidths=[2.0 * inch, 0.8 * inch, 3.2 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2C3E50")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F9FA")]),
            ]
        )
    )
    story.extend([table, Spacer(1, 0.15 * inch)])

    for heading, bullets in SECTIONS:
        story.append(Paragraph(heading, styles["Heading1"]))
        for bullet in bullets:
            story.append(Paragraph("- " + bullet, styles["Normal"]))
        story.append(Spacer(1, 0.08 * inch))

    story.append(PageBreak())
    story.append(Paragraph("Simple Code Examples", styles["Heading1"]))
    for heading, code in CODE_BLOCKS:
        story.append(Paragraph(heading, styles["Heading2"]))
        story.append(Preformatted(code, styles["SmallCode"]))
        story.append(Spacer(1, 0.1 * inch))

    story.append(Paragraph("Learning Path", styles["Heading1"]))
    for index, item in enumerate(
        [
            "Understand the attendance flow from faculty to student.",
            "Learn how JSON moves through REST endpoints.",
            "Learn Hibernate entities and database tables.",
            "Learn Docker Compose by reading one service block at a time.",
        ],
        start=1,
    ):
        story.append(Paragraph(f"{index}. {item}", styles["Normal"]))

    pdf = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )
    pdf.build(story)


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)
